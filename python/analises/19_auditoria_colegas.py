# -*- coding: utf-8 -*-
"""Auditoria: toda afirmação do COMPARACAO_COLEGAS.md contra a fonte."""

import re

import pandas as pd
from docx import Document

from banking.dados import DIR_BRUTOS, carregar_processada

COL = DIR_BRUTOS / "colegas"
ok = lambda b: "OK " if b else "!! "


def cabecalho(t):
    print("\n" + "=" * 78)
    print(t)
    print("=" * 78)


# ---------------------------------------------------------------- 1
cabecalho("1. A TABELA DO DENI: o CSV bate com o documento?")

csv = pd.read_csv(COL / "deni" / "submissao_politica.csv", sep=";")
doc = Document(COL / "deni" / "documento_politica.docx")
tab_doc = doc.tables[1]

def num(txt):
    """Aceita 1.286,20 / 1,286.20 / 87,00 / 0.74 — os dois formatos aparecem."""
    t = re.sub(r"[^\d.,]", "", str(txt))
    if not t:
        return None
    if "," in t and "." in t:
        # o ultimo separador e o decimal
        dec = max(t.rfind(","), t.rfind("."))
        t = re.sub(r"[.,]", "", t[:dec]) + "." + t[dec + 1:]
    else:
        t = t.replace(",", ".")
        if t.count(".") > 1:  # 1.286.20
            i = t.rfind(".")
            t = t[:i].replace(".", "") + "." + t[i + 1:]
    try:
        return float(t)
    except ValueError:
        return None

print(f"\n{'score':>5} | {'faixa no CSV':>16} | {'faixa no DOCX':>16} | "
      f"{'perda CSV':>10} | {'perda DOCX':>10} | {'ROI CSV':>8} | {'ROI DOCX':>8}")
print("-" * 78)
divergem_faixa = divergem_perda = divergem_roi = 0
for i in range(1, 11):
    linha_doc = tab_doc.rows[i]
    score = linha_doc.cells[0].text.strip()
    faixa_doc = linha_doc.cells[1].text.strip()
    perda_doc = linha_doc.cells[6].text.strip()
    roi_doc = linha_doc.cells[7].text.strip()
    lc = csv[csv["Score"].astype(str) == score]
    if lc.empty:
        continue
    lc = lc.iloc[0]
    faixa_csv = f"{lc['Faixa de PD Min']} - {lc['Faixa de PD Max']}"
    perda_csv = str(lc["Perda esperada (R$)"]).strip()
    roi_csv = str(lc["ROI esperado"]).strip()
    print(f"{score:>5} | {faixa_csv:>16} | {faixa_doc:>16} | {perda_csv:>10} | "
          f"{perda_doc:>10} | {roi_csv:>8} | {roi_doc:>8}")
    if num(lc["Faixa de PD Max"]) != num(faixa_doc.split("-")[-1]):
        divergem_faixa += 1
    if num(perda_csv) != num(perda_doc):
        divergem_perda += 1
    if num(roi_csv) != num(roi_doc):
        divergem_roi += 1

print(f"\n  {ok(divergem_faixa == 0)}faixas de PD divergentes : {divergem_faixa} de 10")
print(f"  {ok(divergem_perda == 0)}perdas divergentes       : {divergem_perda} de 10")
print(f"  {ok(divergem_roi == 0)}ROIs divergentes         : {divergem_roi} de 10")

print("\n  Alavancas (taxa, prazo, entrada) — essas eu usei na reconstrucao:")
iguais = True
for i in range(1, 11):
    ld = tab_doc.rows[i]
    score = ld.cells[0].text.strip()
    lc = csv[csv["Score"].astype(str) == score]
    if lc.empty:
        continue
    lc = lc.iloc[0]
    for col_csv, idx_doc, nome in (("Taxa a.m.", 3, "taxa"), ("Prazo máx.", 4, "prazo"),
                                   ("Entrada mín.", 5, "entrada")):
        a, b = str(lc[col_csv]).strip(), ld.cells[idx_doc].text.strip()
        if num(a) != num(b):
            print(f"    !! score {score} {nome}: CSV={a} DOCX={b}")
            iguais = False
print(f"  {ok(iguais)}taxa, prazo e entrada batem entre CSV e DOCX")

print("\n  O ROI por faixa dele e crescente ou decrescente no risco?")
rois_csv = [num(csv[csv['Score'] == s]['ROI esperado'].iloc[0]) for s in range(10, 3, -1)]
rois_doc = [num(tab_doc.rows[11 - s].cells[7].text) for s in range(10, 3, -1)]
print(f"    CSV  (score 10->4): {rois_csv}  -> {'CRESCE' if rois_csv[-1] > rois_csv[0] else 'cai'}")
print(f"    DOCX (score 10->4): {rois_doc}  -> {'cresce' if rois_doc[-1] > rois_doc[0] else 'CAI'}")

# ---------------------------------------------------------------- 2
cabecalho("2. O CORTE DE PD: qual e o certo?")
print("\n  O texto do resumo executivo diz 'ponto de corte de PD em 0.11'.")
print(f"  A ultima faixa aprovada (score 4) termina em:")
print(f"    CSV : {csv[csv['Score'] == 4]['Faixa de PD Max'].iloc[0]}")
print(f"    DOCX: {tab_doc.rows[7].cells[1].text.strip()}")
print("\n  A analise usou 0.11 (do texto). Com o CSV seria 0.0941.")

base_c = carregar_processada("C")
from banking.modelo import treinar_modelo_final
modelo = treinar_modelo_final(carregar_processada("A"))
from banking.dados import preparar_base_c
p = preparar_base_c(base_c).reset_index(drop=True)
p["pd"] = modelo.predict_proba(p)[:, 1]
for corte in (0.0941, 0.11, 0.112):
    print(f"    corte {corte:.4f} -> aprovaria {(p['pd'] <= corte).mean():.1%} das propostas do Leo")

# ---------------------------------------------------------------- 3
cabecalho("3. NUMEROS QUE O DOCUMENTO AFIRMA vs FONTE")

tab_ind_deni = doc.tables[3]
print("\n--- os indicadores que o Deni reporta (tabela 3 do doc dele) ---")
for r in tab_ind_deni.rows[1:]:
    print(f"    {r.cells[0].text.strip():<28} {r.cells[1].text.strip()}")

doc_m = Document(COL / "marcelo" / "documento_politica.docx")
print("\n--- os do Marcelo ---")
for r in doc_m.tables[3].rows[1:]:
    print(f"    {r.cells[0].text.strip():<28} {r.cells[1].text.strip()[:58]}")

# ---------------------------------------------------------------- 4
cabecalho("4. FORMATO — reconferindo")
base_b = carregar_processada("B")
ids_b, ids_c = set(base_b["id_contrato"]), set(base_c["id_proposta"])
for quem, sep_m, sep_p in (("deni", ";", ";"), ("marcelo", ",", ",")):
    dm = pd.read_csv(COL / quem / "submissao_modelo.csv", sep=sep_m)
    dp = pd.read_csv(COL / quem / "submissao_politica.csv", sep=sep_p)
    col_id_m = "id_contrato" if "id_contrato" in dm.columns else dm.columns[0]
    col_id_p = "id_proposta" if "id_proposta" in dp.columns else dp.columns[0]
    cm = len(set(dm[col_id_m]) & ids_b)
    cp = len(set(dp[col_id_p]) & ids_c) if col_id_p in dp.columns else 0
    print(f"\n  {quem.upper()}")
    print(f"    modelo  : {len(dm):>5,} linhas · {cm:>5,}/3.000 ids cruzam")
    print(f"    politica: {len(dp):>5,} linhas · {cp:>5,}/5.000 ids cruzam")
    print(f"    exemplo de id do modelo: {dm[col_id_m].iloc[0]}")

# ---------------------------------------------------------------- 5
cabecalho("5. PDs DO MARCELO: base B e base C sao consistentes?")
mb = pd.read_csv(COL / "marcelo" / "submissao_modelo.csv")
mc = pd.read_csv(COL / "marcelo" / "submissao_politica.csv")
print(f"\n  base B: {len(mb):,} linhas · PD media {mb['pd'].mean():.2%} · "
      f"mediana {mb['pd'].median():.2%} · max {mb['pd'].max():.2%}")
print(f"  base C: {len(mc):,} linhas · PD media {mc['pd'].mean():.2%} · "
      f"mediana {mc['pd'].median():.2%} · max {mc['pd'].max():.2%}")
aprov_m = mc["decisao"].eq("APROVAR").mean()
print(f"\n  aprovacao no CSV dele: {aprov_m:.1%}  (ele reporta 49,2% no documento)")
print(f"  {ok(abs(aprov_m - 0.492) < 0.005)}bate com o documento dele")
apr = mc[mc["decisao"] == "APROVAR"]
print(f"  taxa no CSV: min {apr['taxa_am'].min():.4f} · max {apr['taxa_am'].max():.4f} · "
      f"media {apr['taxa_am'].mean():.4f}")
print(f"    (o documento diz 1,85% a 2,18%, media ~1,95%)")
print(f"  prazos ofertados: {sorted(apr['prazo_meses'].unique())}")
print(f"  entradas ofertadas: {sorted(apr['pct_entrada_minima'].unique())}")
print(f"  score minimo aprovado: {apr['score_1a10'].min()}  (documento diz 'scores 2 a 10')")
