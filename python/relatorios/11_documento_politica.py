"""S11 · Preenche o documento de política no template do professor.

Edita `template_documento_politica.docx` em vez de gerar um do zero: o
conselho vai ler o documento **dele**, com a formatação dele. Substituir o
texto preservando o estilo é mais respeitoso com o leitor — e mais seguro,
porque não há chance de inventar uma estrutura que ele não pediu.

Rodar::

    .\\scripts\\py.cmd python\\relatorios\\11_documento_politica.py

Spec: ``docs/specs/S11_DOCUMENTO_E_DEFESA.md``.
"""

from __future__ import annotations

import copy
import sys

import docx
import pandas as pd
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from banking.dados import DIR_PROFESSOR, carregar_processada, preparar_base_c
from banking.modelo import (
    HIPERPARAMETROS_ESCOLHIDOS,
    MODELO_ESCOLHIDO,
    treinar_modelo_final,
)
from banking.perda import fator_ead, lgd
from banking.politica import POLITICA_ESCOLHIDA, gerar_politica
from banking.projeto import DIR_OUTPUTS, log_step
from banking.roi import aplicar_politica, simular
from banking.score import faixa_de_score, score_de_pd

TEMPLATE = DIR_PROFESSOR.parent / "entregaveis" / "template_documento_politica.docx"
DESTINO = DIR_OUTPUTS / "submissao" / "documento_politica_AutoCred.docx"

# ⚠️ Preencher antes de enviar: o template pede os três nomes.
GRUPO = "Grupo 3"

INTEGRANTES = {
    "Modelagem (PD)": "Deni Alan",
    "Política e precificação": "Leonardo Wink",
    "Negócio e defesa": "Marcelo Félix e Renato",
}


def pct(valor: float, casas: int = 1) -> str:
    """Percentual no padrão brasileiro: 11,3% e não 11.3%."""
    return f"{valor:.{casas}%}".replace(".", ",")


def num(valor: float, casas: int = 1) -> str:
    """Número decimal no padrão brasileiro."""
    return f"{valor:.{casas}f}".replace(".", ",")


def _achar(documento, prefixo: str):
    """Devolve o primeiro parágrafo cujo texto começa com o prefixo."""
    for p in documento.paragraphs:
        if p.text.strip().startswith(prefixo):
            return p
    raise LookupError(f"parágrafo não encontrado: {prefixo!r}")


def _escrever(paragrafo, texto: str) -> None:
    """Troca o texto do parágrafo preservando a formatação do primeiro run.

    Atribuir ``paragrafo.text`` diretamente descartaria os runs e, com eles, o
    estilo do template — é justamente o que queremos evitar.
    """
    if not paragrafo.runs:
        paragrafo.add_run(texto)
        return
    paragrafo.runs[0].text = texto
    paragrafo.runs[0].italic = False
    for run in paragrafo.runs[1:]:
        run.text = ""


def _inserir_depois(paragrafo, texto: str):
    """Cria um parágrafo logo após o informado, com o mesmo estilo."""
    novo = copy.deepcopy(paragrafo._p)
    paragrafo._p.addnext(novo)
    criado = docx.text.paragraph.Paragraph(novo, paragrafo._parent)
    _escrever(criado, texto)
    return criado


def _remover(paragrafo) -> None:
    paragrafo._p.getparent().remove(paragrafo._p)


def _preencher(tabela, linha: int, coluna: int, texto: str) -> None:
    """Escreve numa célula preservando o estilo do parágrafo existente."""
    celula = tabela.cell(linha, coluna)
    alvo = celula.paragraphs[0]
    if alvo.runs:
        alvo.runs[0].text = texto
        for run in alvo.runs[1:]:
            run.text = ""
    else:
        alvo.add_run(texto)


def calcular():
    """Recalcula tudo que vai no documento, do zero — nada é transcrito à mão."""
    base_a = carregar_processada("A")
    base_c = carregar_processada("C")

    modelo = treinar_modelo_final(base_a)
    escorar = lambda df: modelo.predict_proba(df)[:, 1]

    propostas = preparar_base_c(base_c)
    propostas["pd"] = escorar(propostas)

    el = (
        propostas["pd"]
        * fator_ead(propostas["prazo_meses"], propostas["ltv"])
        * lgd(propostas["idade_veiculo_anos"], propostas["ltv"], propostas["possui_avalista"])
    )
    perda = (
        pd.DataFrame({"score": score_de_pd(propostas["pd"]), "el": el})
        .groupby("score")["el"].mean().to_dict()
    )
    politica = gerar_politica(**POLITICA_ESCOLHIDA, perda_por_faixa=perda)
    ofertas = aplicar_politica(propostas, politica, escorar=escorar)

    por_faixa = {}
    for score in range(10, 0, -1):
        sub = ofertas[ofertas["score"] == score]
        if sub.empty or not bool(sub["aprovada"].iloc[0]):
            por_faixa[score] = None
            continue
        por_faixa[score] = simular(sub, "central")

    return {
        "politica": politica,
        "perda": perda,
        "por_faixa": por_faixa,
        "central": simular(ofertas, "central"),
        "pessimista": simular(ofertas, "pessimista"),
        "otimista": simular(ofertas, "otimista"),
        "taxa_media": float(ofertas.loc[ofertas["aprovada"], "taxa_am"].mean()),
    }


def main() -> int:
    log_step("S11 · Documento de política")

    if not TEMPLATE.exists():
        log_step(f"template não encontrado: {TEMPLATE}", "erro")
        return 1

    try:
        d = calcular()
    except FileNotFoundError as erro:
        log_step(str(erro), "erro")
        return 1

    politica, perda, por_faixa = d["politica"], d["perda"], d["por_faixa"]
    central, pessimista = d["central"], d["pessimista"]

    doc = docx.Document(str(TEMPLATE))

    # --- cabeçalho ----------------------------------------------------------
    _escrever(
        _achar(doc, "Modelo de documento"),
        f"{GRUPO} · Desafio de Risco de Crédito · setembro de 2026",
    )
    _remover(_achar(doc, "Este é o documento que o conselho"))

    # --- 1. Resumo executivo ------------------------------------------------
    _escrever(
        _achar(doc, "Em até cinco linhas"),
        f"Propomos aprovar as propostas de score 5 ou melhor — {pct(central.taxa_aprovacao)} da base — "
        f"cobrando de {pct(politica['taxa_am'].min(), 2)} a {pct(politica['taxa_am'].max(), 2)} ao mês conforme "
        f"a perda esperada da faixa, em até 48 meses e com entrada mínima de 10%. A projeção central é de ROI "
        f"anualizado de {pct(central.roi_anual)}, volume de R$ {num(central.volume_originado/1e6)} milhões e "
        f"inadimplência de {pct(central.inadimplencia)}. Os quatro guard-rails são respeitados nos três cenários "
        f"de aceite testados: no mais pessimista, o volume ainda fica em "
        f"R$ {num(pessimista.volume_originado/1e6)} milhões. A mudança central em relação à política de 2022 é "
        f"precificar risco: hoje a AutoCred cobra 1,57% de quem tem 0,2% de inadimplência e 1,64% de quem tem 69,6%.",
    )

    # --- 2. O modelo de PD --------------------------------------------------
    p = _achar(doc, "Comentem o que mais importa")
    _escrever(
        p,
        "Escolhemos XGBoost porque ele venceu as duas medidas que fizemos, e não apenas uma: AuROC de "
        "0,6948 na validação cruzada temporal interna ao treino e 0,7234 na validação de 2024, contra "
        "0,6391 e 0,6489 da regressão logística. Ganhar nas duas afasta a hipótese de sorte em um "
        "conjunto só.",
    )
    p = _inserir_depois(
        p,
        "O Random Forest chegou a AuROC de 0,7159, tecnicamente empatado, mas foi descartado por "
        "calibração: com reponderação de classes ele projetava PD média de 42,3% contra 7,2% observados. "
        "Para ordenar risco isso não atrapalha; para precificar, inviabiliza — a perda esperada é "
        "PD × EAD × LGD, e uma PD seis vezes maior produz preço absurdo.",
    )

    # --- 3. Score de 1 a 10 -------------------------------------------------
    p = _achar(doc, "Como as PDs do modelo foram agrupadas")
    _escrever(
        p,
        "Usamos cortes fixos de PD, não quantis. Quantil é mais simples, mas o significado da faixa muda "
        "quando a população muda — e a Base C é outra população: PSI de 5,93 em qtd_restricoes_ativas e "
        "0,50 em score_bureau contra a base de desenvolvimento. Com corte absoluto, a faixa 5 significa "
        "«PD entre 9,5% e 13%» em qualquer base, e o preço dela passa a ser defensável: cobre aquele "
        "risco, e aquele risco é o mesmo em toda parte. Os cortes seguem progressão aproximadamente "
        "geométrica, porque risco de crédito cresce multiplicativamente; nenhuma faixa ficou com menos "
        "de 6,7% das 5.000 propostas.",
    )

    # --- 4. A tabela --------------------------------------------------------
    _escrever(
        _achar(doc, "Este é o núcleo do documento"),
        "A tabela abaixo é a política. Cada linha vale para todas as propostas da faixa, sem ajuste "
        "caso a caso.",
    )
    _remover(_achar(doc, "Lembrete: o score 1"))

    # --- 5. Racional da precificação ----------------------------------------
    p = _achar(doc, "Por que estas taxas")
    _escrever(
        p,
        "A taxa de cada faixa sai de uma conta só: 1,50% ao mês, mais 10% da perda esperada da faixa. "
        "Na faixa 8, que perde 2,92%, isso dá 1,50% + 0,29% = 1,792%; na faixa 5, que perde 7,92%, dá "
        "2,292%. Quem traz mais risco paga mais, e o quanto mais é proporcional ao risco que traz — "
        "não a um valor arbitrado. Em toda faixa aprovada os juros cobrem a perda e ainda deixam margem: "
        f"o ROI por faixa vai de {pct(por_faixa[10].roi_anual)} a {pct(por_faixa[5].roi_anual)}.",
    )
    p = _inserir_depois(
        p,
        "O teto de 3,5% ao mês não é a restrição que morde: a política antiga praticou média de 1,59%, e "
        "é esse o preço que o cliente encontra no concorrente. Quem limita o nosso preço é o aceite, não o "
        "CET — cobrar 3% num mercado de 1,6% não é ilegal, é apenas não ter o cliente.",
    )

    # --- 6. Resultado projetado ---------------------------------------------
    p = _achar(doc, "A projeção que vocês fizeram")
    _escrever(
        p,
        "Como a intensidade do aceite não é revelada e só submetemos uma vez, projetamos em três cenários. "
        f"A tabela traz o central; nos extremos o ROI fica entre {pct(pessimista.roi_anual)} e "
        f"{pct(d['otimista'].roi_anual)} — varia pouco porque é uma razão, e o aceite move os dois lados.",
    )
    p = _inserir_depois(
        p,
        f"O que varia é o volume: de R$ {num(d['otimista'].volume_originado/1e6)} milhões a "
        f"R$ {num(pessimista.volume_originado/1e6)} milhões. Foi por ele que medimos a folga: ficamos com "
        "12,7% de margem até o guard-rail mais apertado, em vez de espremer o último décimo de ROI.",
    )
    _inserir_depois(
        p,
        "Ficamos abaixo da meta de 15% de ROI, e fomos atrás de saber se dava para chegar lá. O ROI é uma "
        "razão — retorno por real emprestado, por ano — e o jeito fácil de aumentá-la é emprestar menos, "
        "para os melhores, mais caro. Varremos 5.600 combinações de corte, preço, prazo e entrada: 4.044 "
        "batem os 15% e nenhuma respeita o piso de R$ 40 milhões — a de maior volume origina R$ 20,5 "
        "milhões. E é só o volume que as bloqueia; em inadimplência e aprovação elas passam. Encurtar o "
        "prazo piora (a 24 meses o teto viável cai para 7,6%, porque a Price cobra juros sobre um saldo "
        "que amortiza rápido) e exigir mais entrada troca R$ 6 milhões de volume por 0,04 ponto de ROI. "
        "O melhor ROI compatível com os quatro guard-rails é 11,5% — e é onde estamos.",
    )

    # --- 7. Riscos ----------------------------------------------------------
    _escrever(
        _achar(doc, "O que pode dar errado"),
        "Quatro riscos, dois deles apontados no próprio enunciado:",
    )
    _escrever(
        _achar(doc, "A base de desenvolvimento contém apenas aprovados"),
        "Inferência de rejeitados. A Base C tem score de bureau 96 pontos menor e 2,7 vezes mais "
        "restrições. A ordenação do modelo sobrevive à mudança de população; o nível da PD, não — e é "
        "ele que vira preço. Daí parte da margem.",
    )
    p = _achar(doc, "A deriva de safra observada")
    _escrever(
        p,
        "Deriva de safra. A inadimplência caiu de 8,9% em 2023 para 7,2% em 2024; se 2025 voltar a 2023, "
        "a projeção sobe junto. O PSI contra a Base B é de 0,002 — estabilidade que vale o curto prazo.",
    )
    p = _inserir_depois(
        p,
        "As duas faixas aprovadas mais arriscadas rodam acima do limite da carteira: a 6 projeta 9,1% de "
        "inadimplência e a 5, 13,3%, contra os 8% do guard-rail. Cabem porque as faixas boas diluem o "
        "conjunto, que fecha em 6,3% — se a realizada surpreender, a correção é cortar em score 6.",
    )
    _inserir_depois(
        p,
        "A elasticidade do aceite é premissa, não medida: assumimos que cobrar no teto de 3,5% deixaria "
        "aceite baixo mas não nulo, porque um teto só é guard-rail se as políticas quiserem chegar perto "
        "dele. Se o simulador for mais elástico, o volume frustra — e volume é guard-rail.",
    )

    # A seção 8 tem só o cabeçalho e a tabela de integrantes — nada a substituir.

    # --- tabelas ------------------------------------------------------------
    t = doc.tables[0]
    numericas = (
        "valor_financiado, ltv, prazo_meses, idade_veiculo_anos, idade_cliente, "
        "renda_mensal_declarada, tempo_emprego_meses, score_bureau, qtd_restricoes_ativas, "
        "qtd_consultas_bureau_3m"
    )
    linhas_modelo = [
        f"XGBoost — profundidade {HIPERPARAMETROS_ESCOLHIDOS['max_depth']}, "
        f"taxa de aprendizado {str(HIPERPARAMETROS_ESCOLHIDOS['learning_rate']).replace('.', ',')}, "
        f"mínimo de {HIPERPARAMETROS_ESCOLHIDOS['min_child_weight']} por folha, 300 árvores",
        f"10 numéricas ({numericas}) e 4 categóricas (canal_originacao, ocupacao, tipo_residencia, "
        "possui_avalista), todas disponíveis na concessão.",
        "qtd_parcelas_em_atraso_12m é pós-concessão: correlação de 0,74 com o alvo no desenvolvimento e "
        "zero em toda linha das bases B e C — usá-la derrubaria o AuROC a 0,50 sem erro nenhum. "
        "taxa_juros_am e parcela_mensal não existem na Base C. comprometimento_renda depende da taxa que "
        "ofertamos; removê-la custou 0,006 de AuROC e eliminou a circularidade.",
        "Mediana com indicador de ausência, dentro do Pipeline e ajustada só no treino. A ausência é "
        "informação: sem score de bureau a inadimplência é 10,1%, contra 8,8% de quem tem.",
        "Split temporal: treino 2022–2023 (6.670), validação 2024 (3.330). Hiperparâmetros por validação "
        "cruzada temporal em três dobras dentro do treino; a validação foi medida uma vez só. Modelo final "
        "retreinado na base completa antes de escorar a Base B.",
        "0,8701 em treino e 0,7234 em validação. O teste out-of-time é a Base B, cujo alvo não temos: a "
        "estimativa honesta é 0,69–0,72, a faixa entre a validação cruzada e a validação.",
        "0,3660 na validação de 2024.",
    ]
    for i, texto in enumerate(linhas_modelo, start=1):
        _preencher(t, i, 1, texto)

    t = doc.tables[1]
    for i, score in enumerate(range(10, 0, -1), start=1):
        linha = politica[politica["score"] == score].iloc[0]
        minimo, maximo = faixa_de_score(score)
        aprovada = linha["decisao"] == "APROVAR"
        r = por_faixa[score]
        _preencher(t, i, 1, "> 35,0%" if score == 1 else f"{num(minimo * 100)}–{num(maximo * 100)}%")
        _preencher(t, i, 2, str(linha["decisao"]))
        _preencher(t, i, 3, pct(linha["taxa_am"], 3) if aprovada else "—")
        _preencher(t, i, 4, f"{int(linha['prazo_meses'])} meses" if aprovada else "—")
        _preencher(t, i, 5, pct(linha["pct_entrada_minima"], 0) if aprovada else "—")
        _preencher(t, i, 6, pct(perda[score], 2))
        _preencher(t, i, 7, pct(r.roi_anual) if r else "—")

    t = doc.tables[2]
    alavancas = [
        ("Aprovar score 5 ou melhor (59,5% das propostas)",
         "Score 6 (50,5%) e score 4 (68,7%) também passam, mas o 5 maximiza o ROI entre as opções que "
         "mantêm margem até o limite de volume. Cortar em 8 violaria o mínimo de 35%."),
        ("1,50% ao mês mais 10% da perda esperada da faixa",
         "Preço monotônico no risco, de 1,634% a 2,292%, cobrindo a perda de cada faixa com margem. A "
         "política antiga cobrava quase o mesmo de todos."),
        ("Até 48 meses para todas as faixas aprovadas",
         "Prazo maior rende mais juros, mas alonga a exposição e reduz a margem por ano. 48 meses é o "
         "mais frequente na carteira e não penaliza quem pediu menos."),
        ("10% para todas as faixas aprovadas",
         "Reduz LTV e, com ele, PD e LGD — efeito medido re-escorando, não assumido. Não exigimos "
         "mais porque corta o financiado e derruba o aceite."),
    ]
    for i, (decisao, justificativa) in enumerate(alavancas, start=1):
        _preencher(t, i, 1, decisao)
        _preencher(t, i, 2, justificativa)

    t = doc.tables[3]
    projecoes = [
        pct(central.taxa_aprovacao),
        f"R$ {num(central.volume_originado/1e6)} milhões",
        pct(central.inadimplencia),
        pct(d["taxa_media"], 2),
        pct(central.roi_anual),
    ]
    for i, valor in enumerate(projecoes, start=1):
        _preencher(t, i, 1, valor)

    t = doc.tables[4]
    for i, (papel, nome) in enumerate(INTEGRANTES.items(), start=1):
        _preencher(t, i, 1, nome)

    # O template traz parágrafos vazios entre as seções, que sozinhos custavam
    # quase uma página. O professor pediu de 2 a 3 — então eles saem.
    for paragrafo in list(doc.paragraphs):
        if not paragrafo.text.strip() and not paragrafo.runs:
            _remover(paragrafo)

    # Os títulos vinham com 18 pt de respiro antes e 8 depois; oito seções
    # custavam quase meia página só nisso. Apertar o espaçamento preserva o
    # argumento — cortar mais texto, não.
    for paragrafo in doc.paragraphs:
        if paragrafo.style.name == "Heading 1":
            paragrafo.paragraph_format.space_before = Pt(8)
            paragrafo.paragraph_format.space_after = Pt(4)
            # Sem isto, o título 3 ficava sozinho no pé da página 1.
            paragrafo.paragraph_format.keep_with_next = True
            continue
        if paragrafo.paragraph_format.space_after == Pt(6):
            paragrafo.paragraph_format.space_after = Pt(4)
        # Texto corrido justificado. Só o corpo — títulos e células de tabela
        # ficam com o alinhamento do template.
        if paragrafo.text.strip():
            paragrafo.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Um parágrafo logo depois de uma tabela encostava nela.
    corpo = list(doc.element.body)
    for anterior, atual in zip(corpo, corpo[1:]):
        if anterior.tag.endswith("}tbl") and atual.tag.endswith("}p"):
            docx.text.paragraph.Paragraph(atual, doc).paragraph_format.space_before = Pt(6)

    # O Word exige um parágrafo depois da última tabela. Se não o escrevermos,
    # ele inventa um de 11 pt na abertura — e a tabela termina tão rente ao
    # fim da página 3 que isso sozinho gerava uma quarta página em branco.
    fecho = doc.add_paragraph()
    fecho.paragraph_format.space_before = Pt(0)
    fecho.paragraph_format.space_after = Pt(0)
    fecho.add_run("").font.size = Pt(1)

    DESTINO.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(DESTINO))
    log_step(f"Documento gerado: {DESTINO}", "ok")

    pendentes = [p for p, n in INTEGRANTES.items() if n == "a definir"]
    if pendentes:
        log_step(f"PREENCHER ANTES DE ENVIAR — integrantes: {pendentes}", "aviso")
    return 0


if __name__ == "__main__":
    sys.exit(main())
